# Round-trip test: generate workbook → simulate human entry → parse back →
# merge into a temp price-log.csv → assert dedupe + per-stem normalization.
# Run with `python3 -m pytest tests/` or plain `python3 tests/test_sheet_io.py`
# from afuvai-price-monitor/ (or anywhere — paths are resolved from this file).

import shutil
import sys
import tempfile
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "src"))

from openpyxl import load_workbook

from pricemonitor.models import PRICE_LOG_COLUMNS, ProposalRow
from pricemonitor.sheet_io import (
    QUICK_ENTRY_COLUMNS,
    TAB_PROPOSALS,
    TAB_QUICK_ENTRY,
    build_workbook,
    merge_rows_into_price_log,
    parse_proposals_tab,
    parse_quick_entry_tab,
    read_price_log,
    write_proposals,
)

EXPECTED_TABS = [
    "Read Me", "Quick Entry", "Price Log", "Proposals", "Vendors",
    "Substitutions & Units", "Tier Recipes", "Seasonal Calendar",
    "Comparison", "Digest Archive",
]


def _make_workspace() -> tuple[Path, Path, Path]:
    workdir = Path(tempfile.mkdtemp(prefix="afuvai-sheet-io-"))
    data_dir = workdir / "data"
    data_dir.mkdir()
    for name in ("price-log.csv", "proposals.csv", "tier-price-history.csv"):
        shutil.copy(PROJECT_ROOT / "data" / name, data_dir / name)
    workbook_path = workdir / "AFUVAI Vendor Price Monitor.xlsx"
    return workdir, data_dir, workbook_path


def _append_quick_entry(workbook_path: Path, entries: list[dict]) -> None:
    # Simulate Brian typing rows into Quick Entry (row 1 note, row 2 headers).
    workbook = load_workbook(workbook_path)
    sheet = workbook[TAB_QUICK_ENTRY]
    next_row = max(sheet.max_row + 1, 3)
    for entry in entries:
        for column, header in enumerate(QUICK_ENTRY_COLUMNS, start=1):
            sheet.cell(row=next_row, column=column, value=entry.get(header, ""))
        next_row += 1
    workbook.save(workbook_path)


def test_workbook_has_all_tabs():
    _, data_dir, workbook_path = _make_workspace()
    build_workbook(workbook_path, data_dir=data_dir)
    workbook = load_workbook(workbook_path)
    assert workbook.sheetnames == EXPECTED_TABS, workbook.sheetnames
    # Price Log tab mirrors the CSV: note row + header + 2 sample rows.
    log_sheet = workbook["Price Log"]
    headers = [cell.value for cell in log_sheet[2]]
    assert headers == PRICE_LOG_COLUMNS
    assert log_sheet.max_row == 4  # note + header + 2 sample rows


def test_quick_entry_round_trip_with_normalization_and_dedupe():
    _, data_dir, workbook_path = _make_workspace()
    build_workbook(workbook_path, data_dir=data_dir)

    bunch_quote = {
        # No units_per_quote_unit → falls back to standard_rose default of 25.
        "date": "2026-07-16", "vendor_id": "mayesh-wholesale-las-vegas",
        "stem_id": "standard_rose", "variety": "Playa Blanca", "grade": "standard",
        "quoted_price": 22.50, "quote_unit": "bunch", "units_per_quote_unit": "",
        "currency": "USD", "source": "phone", "baseline": "true",
        "valid_until": "2026-08-16", "disruption_tag": "", "notes": "first quote",
        "entered_by": "brian",
    }
    explicit_units_quote = {
        # Explicit per-quote override: hydrangea bunch of 10 (default is 5).
        "date": "2026-07-16", "vendor_id": "greenleaf-wholesale-sparks",
        "stem_id": "hydrangea", "variety": "", "grade": "standard",
        "quoted_price": 30.00, "quote_unit": "bunch", "units_per_quote_unit": 10,
        "currency": "USD", "source": "email", "baseline": "true",
        "valid_until": "", "disruption_tag": "", "notes": "", "entered_by": "brian",
    }
    per_stem_quote = {
        "date": "2026-07-16", "vendor_id": "las-vegas-flower-market",
        "stem_id": "peony", "variety": "Sarah Bernhardt", "grade": "premium",
        "quoted_price": 4.25, "quote_unit": "stem", "units_per_quote_unit": 1,
        "currency": "USD", "source": "manual", "baseline": "true",
        "valid_until": "", "disruption_tag": "", "notes": "", "entered_by": "brian",
    }
    duplicate_of_sample = {
        # Matches sample-0001's dedupe key (vendor, stem, date, per-stem 0.75).
        "date": "2026-07-15", "vendor_id": "las-vegas-flower-market",
        "stem_id": "standard_rose", "variety": "Freedom Red", "grade": "standard",
        "quoted_price": 18.75, "quote_unit": "bunch", "units_per_quote_unit": 25,
        "currency": "USD", "source": "manual", "baseline": "false",
        "valid_until": "", "disruption_tag": "", "notes": "re-keyed by accident",
        "entered_by": "brian",
    }
    _append_quick_entry(
        workbook_path,
        [bunch_quote, explicit_units_quote, per_stem_quote, duplicate_of_sample, duplicate_of_sample],
    )

    parsed = parse_quick_entry_tab(workbook_path)
    assert len(parsed) == 5

    by_stem = {row.stem_id: row for row in parsed[:3]}
    assert by_stem["standard_rose"].price_per_stem == 0.9  # 22.50 / default 25
    assert by_stem["standard_rose"].units_per_quote_unit == 25  # default recorded back
    assert "defaulted" in by_stem["standard_rose"].notes
    assert by_stem["hydrangea"].price_per_stem == 3.0  # 30 / explicit 10, not default 5
    assert by_stem["peony"].price_per_stem == 4.25
    assert by_stem["standard_rose"].baseline is True

    log_path = data_dir / "price-log.csv"
    added, skipped = merge_rows_into_price_log(parsed, csv_path=log_path)
    # 3 new rows added; the sample duplicate + its own repeat both skipped.
    assert len(added) == 3, [row.dedupe_key() for row in added]
    assert len(skipped) == 2

    merged = read_price_log(log_path)
    assert len(merged) == 5  # 2 samples + 3 new
    keys = [row.dedupe_key() for row in merged]
    assert len(keys) == len(set(keys)), "dedupe keys must be unique after merge"
    for row in added:
        assert row.id.startswith("pl-")
        assert row.entered_at
        assert row.status == "confirmed"

    # Re-merging the same parse is a no-op (idempotent weekly job).
    added_again, skipped_again = merge_rows_into_price_log(parsed, csv_path=log_path)
    assert added_again == []
    assert len(skipped_again) == 5
    assert len(read_price_log(log_path)) == 5


def test_proposals_round_trip_and_confirm_flow():
    _, data_dir, workbook_path = _make_workspace()
    proposal = ProposalRow(
        id="prop-0001", date="2026-07-14", vendor_id="mayesh-wholesale-las-vegas",
        stem_id="dahlia", variety="Cornel Bronze", grade="premium",
        quoted_price=32.00, quote_unit="bunch", units_per_quote_unit=10,
        price_per_stem=3.2, source="email", gmail_thread_id="t-123",
        gmail_message_id="m-456", extracted_at="2026-07-14T10:00:00+00:00",
        confidence=0.92, status="proposed",
    )
    write_proposals([proposal], data_dir / "proposals.csv")
    build_workbook(workbook_path, data_dir=data_dir)

    # Simulate Brian flipping the status column to confirmed in the sheet.
    workbook = load_workbook(workbook_path)
    sheet = workbook[TAB_PROPOSALS]
    headers = [cell.value for cell in sheet[2]]
    status_column = headers.index("status") + 1
    assert sheet.cell(row=3, column=status_column).value == "proposed"
    sheet.cell(row=3, column=status_column, value="confirmed")
    workbook.save(workbook_path)

    proposals = parse_proposals_tab(workbook_path)
    assert len(proposals) == 1
    assert proposals[0].status == "confirmed"
    assert proposals[0].gmail_thread_id == "t-123"
    assert proposals[0].confidence == 0.92

    log_path = data_dir / "price-log.csv"
    added, skipped = merge_rows_into_price_log(
        [proposals[0].to_price_log_row()], csv_path=log_path,
    )
    assert len(added) == 1 and not skipped
    assert added[0].price_per_stem == 3.2
    assert added[0].status == "confirmed"
    assert added[0].source == "email"


if __name__ == "__main__":
    test_workbook_has_all_tabs()
    test_quick_entry_round_trip_with_normalization_and_dedupe()
    test_proposals_round_trip_and_confirm_flow()
    print("OK: all sheet_io round-trip tests passed")
